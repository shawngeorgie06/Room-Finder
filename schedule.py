import csv
from datetime import datetime, time
import re
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

# Above the bundled 7,170-row exports, while bounding compressed XLSX work.
MAX_SCHEDULE_ROWS = 50_000
MAX_SCHEDULE_COLUMNS = 128
MAX_CELL_CHARS = 4096
MAX_XLSX_EXPANDED_BYTES = 64 * 1024 * 1024
MAX_XLSX_MEMBERS = 1000
MAX_DIAGNOSTICS = 10
REQUIRED_COLUMNS = {'Days', 'Times', 'Location', 'Delivery Mode'}
ONLINE_MODES = {'Online Newark', 'Online Virtual', 'Synchronous Online', 'Asynchronous Online', 'Online'}
UNSCHEDULED = {'', 'TBA', 'TBD', 'ARR', 'ARRANGED'}


class ScheduleValidationError(ValueError):
    """Bounded, actionable errors without echoing untrusted cell contents."""

    def __init__(self, diagnostics, error_count=None):
        self.diagnostics = tuple(diagnostics[:MAX_DIAGNOSTICS])
        self.error_count = error_count if error_count is not None else len(diagnostics)
        suffix = f' ({self.error_count} errors total)' if self.error_count > len(self.diagnostics) else ''
        super().__init__('Invalid schedule: ' + '; '.join(self.diagnostics) + suffix)


def _headers(values):
    if len(values) > MAX_SCHEDULE_COLUMNS:
        raise ScheduleValidationError(['column limit exceeded; export only schedule columns'])
    headers = [str(h).strip() if h is not None else '' for h in values]
    nonempty = [h for h in headers if h]
    if len(set(nonempty)) != len(nonempty):
        raise ScheduleValidationError(['duplicate column headers; use a fresh Banner export'])
    missing = REQUIRED_COLUMNS - set(headers)
    if missing:
        raise ScheduleValidationError(['missing required columns: ' + ', '.join(sorted(missing))])
    return headers


def _row(headers, values, number, *, sparse=False):
    if len(values) > MAX_SCHEDULE_COLUMNS:
        raise ScheduleValidationError([f'row {number}: column limit exceeded; export only schedule columns'])
    if not sparse and len(values) < len(headers) and any(v is not None and str(v).strip() for v in values):
        raise ScheduleValidationError([f'row {number}: missing cells; check CSV quoting or export again'])
    if len(values) > len(headers) and any(v is not None and str(v).strip() for v in values[len(headers):]):
        raise ScheduleValidationError([f'row {number}: extra cells; check CSV quoting or export again'])
    if any(v is not None and len(str(v)) > MAX_CELL_CHARS for v in values):
        raise ScheduleValidationError([f'row {number}: cell length limit exceeded; shorten the field'])
    return {h: '' if v is None else str(v) for h, v in zip(headers, values) if h}

DAY_MAP = {"M": 0, "T": 1, "W": 2, "R": 3, "F": 4, "S": 5, "U": 6}
VALID_MODES = {"Face-to-Face", "Hybrid", "Converged Learning"}


def parse_time_range(time_str):
    """Parse '1:00 PM - 2:20 PM' into (time, time). Returns None on failure."""
    if not time_str or not isinstance(time_str, str):
        return None
    parts = time_str.strip().split(" - ")
    if len(parts) != 2:
        return None
    try:
        start = datetime.strptime(parts[0].strip(), "%I:%M %p").time()
        end = datetime.strptime(parts[1].strip(), "%I:%M %p").time()
        return (start, end) if end > start else None
    except ValueError:
        return None


def parse_days(days_str):
    """Parse 'MWF' into [0, 2, 4]. Unknown codes are silently skipped."""
    if not days_str or not isinstance(days_str, str):
        return []
    return [DAY_MAP[ch] for ch in days_str.strip() if ch in DAY_MAP]


def parse_location(location_str):
    """Parse 'KUPF 207' into ('KUPF', '207'). Returns None if invalid."""
    if not location_str or not isinstance(location_str, str):
        return None
    stripped = location_str.strip()
    if not stripped:
        return None
    parts = stripped.split(None, 1)
    if len(parts) < 2:
        return None
    building, room = parts[0], parts[1].strip()
    if any(len(value) > 128 or not re.fullmatch(r'[A-Za-z0-9][A-Za-z0-9 ._/-]*', value)
           for value in (building, room)):
        return None
    return building, room


def _iter_rows_csv(path):
    """Yield logical rows; utf-8-sig accepts CSV exports with or without BOM."""
    with open(path, newline='', encoding='utf-8-sig') as fh:
        reader = csv.reader(fh, strict=True)
        headers = _headers(next(reader, []))
        for number, values in enumerate(reader, 2):
            yield _row(headers, values, number)


def _iter_rows_xlsx(path):
    """Bound ZIP expansion before openpyxl allocates shared strings/styles."""
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > MAX_XLSX_MEMBERS or sum(m.file_size for m in members) > MAX_XLSX_EXPANDED_BYTES:
            raise ScheduleValidationError(['XLSX expanded size limit exceeded (64 MiB / 1000 ZIP members); export CSV instead'])
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        ws = wb.worksheets[0]
        if (ws.max_row and ws.max_row > MAX_SCHEDULE_ROWS + 1) or (ws.max_column and ws.max_column > MAX_SCHEDULE_COLUMNS):
            raise ScheduleValidationError(['XLSX row/column limit exceeded; export a single semester'])
        # Exported dimension metadata may understate the sheet. Stream actual
        # rows instead, so a stale dimension cannot hide malformed meetings.
        ws.reset_dimensions()
        rows = ws.iter_rows(values_only=True)
        headers = _headers(next(rows, []))
        for number, values in enumerate(rows, 2):
            # Excel omits trailing empty cells; their column coordinates still
            # distinguish them from malformed short CSV records.
            yield _row(headers, values, number, sparse=True)
    finally:
        wb.close()


def load_schedule(path, *, report=None):
    """Return a complete, validated single-term list, or raise ValueError.

    Explicit online modes and unscheduled blank/TBA/arranged meetings are excluded:
    they cannot establish physical-room occupancy. Supplied non-placeholder in-person
    fields are still validated, so an exclusion cannot hide a malformed meeting.
    Converged Learning with a physical meeting is included. Blank capacity/term is
    unknown; fractional/negative capacities and mixed known/unknown terms are invalid.
    Limits: 50,000 rows, 128 columns, 4,096 characters/cell, 64 MiB XLSX expansion.
    An optional report dict receives accepted/excluded row counts on success;
    callers using load_schedule(path) keep the original plain-list interface.
    """
    suffix = Path(path).suffix.lower()
    if suffix not in {'.csv', '.xlsx'}:
        raise ScheduleValidationError(['use a CSV or XLSX Banner export'])
    rows = _iter_rows_csv(path) if suffix == '.csv' else _iter_rows_xlsx(path)
    results, diagnostics, terms = [], [], set()
    error_count = 0
    excluded_online = excluded_unscheduled = 0
    try:
        for number, row in enumerate(rows, 2):
            if number > MAX_SCHEDULE_ROWS + 1:
                raise ScheduleValidationError(['schedule row limit exceeded (50,000); export a single semester'])
            if not any(value.strip() for value in row.values()):
                continue
            mode = row.get('Delivery Mode', '').strip()
            if mode in ONLINE_MODES:
                excluded_online += 1
                continue
            values = {field: row.get(field, '').strip() for field in ('Location', 'Days', 'Times', 'Max', 'Term')}
            absent = {field: values[field].upper() in UNSCHEDULED for field in ('Location', 'Days', 'Times')}
            issues = []
            if mode not in VALID_MODES and (mode or not absent['Location']):
                issues.append('Delivery Mode: use a recognized physical or online mode')
            loc = parse_location(values['Location'])
            if not absent['Location'] and loc is None:
                issues.append('Location: provide a building and exact room identifier')
            if not absent['Days'] and (not values['Days'] or any(ch not in DAY_MAP for ch in values['Days'])):
                issues.append('Days: use only MTWRFSU weekday codes')
            parsed = parse_time_range(values['Times'])
            if not absent['Times'] and parsed is None:
                issues.append('Times: use h:mm AM/PM - h:mm AM/PM with end after start on the same day')
            cap = None
            if values['Max']:
                try:
                    decimal = Decimal(values['Max'])
                    if not decimal.is_finite() or decimal < 0 or decimal > 1_000_000 or decimal != decimal.to_integral_value():
                        raise ValueError
                    cap = int(decimal)
                except (InvalidOperation, ValueError):
                    issues.append('Max: provide a nonnegative whole-number enrollment limit or leave blank')
            term = values['Term']
            if term and not re.fullmatch(r'[1-9][0-9]{3}(10|30|50|90)', term):
                issues.append('Term: use a six-digit Banner semester code or leave blank')
            if issues:
                error_count += len(issues)
                diagnostics.extend(f'row {number}: {issue}' for issue in issues[:max(0, MAX_DIAGNOSTICS - len(diagnostics))])
                continue
            if any(absent.values()) or mode not in VALID_MODES:
                excluded_unscheduled += 1
                continue
            terms.add(term)
            if len(terms) > 1:
                raise ScheduleValidationError([f'row {number}: mixed terms; export one semester with a consistent Term column'])
            building, room = loc
            start, end = parsed
            results.append({
                'building': building, 'room': room, 'days': parse_days(values['Days']),
                'time_start': start, 'time_end': end, 'capacity': cap,
                'course': row.get('Course', '').strip(), 'title': row.get('Title', '').strip(),
                'instructor': row.get('Instructor', '').strip(), 'term': term,
            })
    except ScheduleValidationError:
        raise
    except (csv.Error, UnicodeError, zipfile.BadZipFile, KeyError, IndexError, SyntaxError) as error:
        raise ScheduleValidationError(['file structure could not be read; export a fresh CSV or XLSX file']) from error
    finally:
        rows.close()
    if diagnostics:
        raise ScheduleValidationError(diagnostics, error_count)
    if not results:
        raise ScheduleValidationError(['no scheduled in-person entries; online and blank/TBA meetings are excluded'])
    if report is not None:
        report.update(accepted_rows=len(results), excluded_online=excluded_online,
                      excluded_unscheduled=excluded_unscheduled)
    return results


def _time_to_minutes(t):
    return t.hour * 60 + t.minute


def merged_busy_intervals(classes):
    """Return sorted, merged [start, end) minute pairs for one room/day.

    Overlapping, nested, duplicate and touching classes are one busy interval.
    Only calculation boundaries are merged; original course records are untouched.
    """
    intervals = sorted(
        (_time_to_minutes(c["time_start"]), _time_to_minutes(c["time_end"]))
        for c in classes
    )
    merged = []
    for start, end in intervals:
        if merged and start <= merged[-1][1]:
            merged[-1] = (merged[-1][0], max(merged[-1][1], end))
        else:
            merged.append((start, end))
    return merged


def _query_window_minutes(now, min_duration_mins, until):
    """Validate a same-day query and return its minute-resolution boundaries."""
    if not isinstance(now, time):
        raise ValueError("Query time must be a time of day")
    now_min = _time_to_minutes(now)
    if until is not None:
        if not isinstance(until, time):
            raise ValueError("Window end must be a time of day")
        end_min = _time_to_minutes(until)
        if end_min <= now_min:
            raise ValueError("Window end must be after its start on the same day")
    else:
        if (not isinstance(min_duration_mins, int)
                or isinstance(min_duration_mins, bool) or min_duration_mins < 0):
            raise ValueError("Duration must be a nonnegative integer number of minutes")
        end_min = now_min + min_duration_mins
        if end_min > 24 * 60:
            raise ValueError("Duration must not cross midnight")
    return now_min, end_min


def _availability_facts(intervals, now_min, requested_end_min):
    """Compute independent facts from validated query and merged busy intervals."""
    empty = True
    available_from = now_min
    available_until = None
    for start, end in intervals:
        if end <= now_min:
            continue
        if start <= now_min:
            empty = False
            available_from = end
        else:
            available_until = start
            break

    return {
        "empty": empty,
        "matches_requested_window": empty and (
            available_until is None or requested_end_min <= available_until
        ),
        "available_from_min": available_from,
        "available_until_min": available_until,
        "minutes_until_next": (
            available_until - now_min if empty and available_until is not None else None
        ),
    }


def get_availability_facts(classes, now, min_duration_mins=0, until=None):
    """Return availability facts for one room's classes on the queried day.

    `empty` describes the instant, independently of requested-window matching.
    Availability starts now when free, or at the end of the current merged busy
    interval when occupied. Its end is the next merged busy interval's start;
    None means no further scheduled class, not guaranteed access/opening hours.
    `minutes_until_next` describes only the CURRENT free interval, so is None
    when occupied as well as when there is no later class.

    An explicit `until` takes precedence over duration. Invalid/equal/reversed
    explicit windows and durations crossing midnight raise ValueError.
    """
    now_min, end_min = _query_window_minutes(now, min_duration_mins, until)
    return _availability_facts(merged_busy_intervals(classes), now_min, end_min)


def get_room_availability(schedule, weekday, now, building=None, min_duration_mins=0, until=None):
    """Return every active-schedule room's facts, sorted by building and room.

    Inventory includes rooms with no classes on the queried day. Capacity is
    the maximum known schedule enrollment limit across ALL days for that room,
    not physical capacity, and never depends on occupancy or window matching.
    Metadata-only rooms cannot create inventory here.
    """
    if not isinstance(weekday, int) or isinstance(weekday, bool) or not 0 <= weekday <= 6:
        raise ValueError("Weekday must be an integer from 0 to 6")
    now_min, end_min = _query_window_minutes(now, min_duration_mins, until)
    all_rooms = {}
    for entry in schedule:
        if building and entry["building"] != building:
            continue
        key = (entry["building"], entry["room"])
        all_rooms.setdefault(key, []).append(entry)

    results = []
    for (bldg, room), entries in sorted(all_rooms.items()):
        today_classes = [e for e in entries if weekday in e["days"]]
        cap = max(
            (e["capacity"] for e in entries if e.get("capacity") is not None),
            default=None,
        )
        results.append({
            "building": bldg,
            "room": room,
            "capacity": cap,
            **_availability_facts(merged_busy_intervals(today_classes), now_min, end_min),
        })
    return results


def get_empty_rooms(schedule, weekday, now, building=None, min_duration_mins=0, until=None):
    """Compatibility wrapper returning exactly rooms matching the requested window.

    Retains building, room, capacity and minutes_until_next, with additive facts.
    Every returned room has empty=True and matches_requested_window=True.
    """
    return [
        row for row in get_room_availability(
            schedule, weekday, now, building=building,
            min_duration_mins=min_duration_mins, until=until,
        )
        if row["matches_requested_window"]
    ]

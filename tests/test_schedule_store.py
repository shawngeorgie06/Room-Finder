"""Publication failure tests exercise the manifest commit point through HTTP."""
import io
import zipfile
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook
from werkzeug.datastructures import FileStorage

import app as app_module
import schedule as parser
from app import create_app
from tests.test_import_provenance import csv_bytes, isolated, upload, info, AUTH


@pytest.mark.parametrize('failure', ['save', 'payload_replace', 'manifest_replace', 'manifest_write', 'snapshot', 'directory_sync'])
def test_failed_publication_preserves_active_and_restart(isolated, monkeypatch, failure):
    import schedule_store
    storage, _, db = isolated
    client = create_app(visit_db_path=str(db)).test_client()
    assert upload(client).status_code == 200
    before = info(client)
    files = {p.name: p.read_bytes() for p in storage.iterdir()}
    original_replace = schedule_store.os.replace
    def failing_replace(src, dst):
        is_manifest = Path(dst).name == schedule_store.MANIFEST_NAME
        if (failure == 'manifest_replace' and is_manifest) or (failure == 'payload_replace' and not is_manifest):
            raise OSError('synthetic write failure')
        return original_replace(src, dst)
    with monkeypatch.context() as m:
        if failure == 'save':
            def fail_save(self, dst, *args, **kwargs):
                Path(dst).write_bytes(b'partial')
                raise OSError('synthetic save failure')
            m.setattr(FileStorage, 'save', fail_save)
        elif failure == 'manifest_write':
            m.setattr(schedule_store.json, 'dump', lambda *args, **kwargs: (_ for _ in ()).throw(OSError('synthetic manifest failure')))
        elif failure == 'snapshot':
            m.setattr(app_module, '_make_snapshot', lambda *args, **kwargs: (_ for _ in ()).throw(ValueError('synthetic snapshot failure')))
        elif failure == 'directory_sync':
            import stat
            original_fsync = schedule_store.os.fsync
            def fail_directory_sync(fd):
                if stat.S_ISDIR(schedule_store.os.fstat(fd).st_mode):
                    raise OSError('synthetic directory sync failure')
                return original_fsync(fd)
            m.setattr(schedule_store.os, 'fsync', fail_directory_sync)
        else:
            m.setattr(schedule_store.os, 'replace', failing_replace)
        response = upload(client, csv_bytes(Location='TEST CHANGED'))
    assert response.status_code == 500
    assert response.headers['Cache-Control'] == 'no-store'
    assert response.headers['Access-Control-Allow-Origin'] == app_module.BANNER_ORIGIN
    assert info(client) == before
    assert {p.name: p.read_bytes() for p in storage.iterdir()} == files
    assert info(create_app(visit_db_path=str(db)).test_client())['schedule_version'] == before['schedule_version']


def test_same_term_csv_then_xlsx_preserves_history_and_exports_original(isolated):
    storage, legacy, db = isolated
    old = legacy / 'uploaded_schedule_202690.csv'
    old.write_bytes(csv_bytes(Location='LEGACY 101'))
    client = create_app(visit_db_path=str(db)).test_client()
    assert upload(client).status_code == 200
    first_files = {p.name: p.read_bytes() for p in storage.iterdir() if p.suffix == '.csv'}
    wb = Workbook()
    sheet = wb.active
    sheet.append(['Term', 'Days', 'Times', 'Location', 'Max', 'Delivery Mode'])
    sheet.append([202690, 'M', '9:00 AM - 10:00 AM', 'TEST LECT 2', 20, 'Face-to-Face'])
    buffer = io.BytesIO()
    wb.save(buffer)
    original = buffer.getvalue()
    assert upload(client, original, 'replacement.xlsx').status_code == 200
    exported = client.get('/api/export-schedule?term=202690', headers=AUTH)
    assert exported.data == original
    assert '.xlsx' in exported.headers['Content-Disposition']
    assert old.read_bytes() == csv_bytes(Location='LEGACY 101')
    for name, body in first_files.items():
        assert (storage / name).read_bytes() == body
    restarted = create_app(visit_db_path=str(db)).test_client()
    assert restarted.get('/api/export-schedule?term=202690', headers=AUTH).data == original
    assert info(restarted)['source_filename'] == 'replacement.xlsx'


def test_manifest_size_limit_cannot_publish_unreadable_state(isolated, monkeypatch):
    import schedule_store
    storage, _, db = isolated
    client = create_app(visit_db_path=str(db)).test_client()
    assert upload(client).status_code == 200
    original = info(client)
    files = {p.name: p.read_bytes() for p in storage.iterdir()}
    monkeypatch.setattr(schedule_store, 'MAX_MANIFEST_BYTES', (storage / schedule_store.MANIFEST_NAME).stat().st_size + 10)
    assert upload(client, csv_bytes(Term='202650'), 'other-semester.csv').status_code == 500
    assert {p.name: p.read_bytes() for p in storage.iterdir()} == files
    assert info(client) == original


def test_unpublished_payload_does_not_shadow_legacy(isolated):
    import schedule_store
    storage, legacy, db = isolated
    storage.mkdir()
    (storage / '.schedule-data-orphan.csv').write_bytes(csv_bytes(Location='ORPHAN 101'))
    (legacy / 'uploaded_schedule_202690.csv').write_bytes(csv_bytes(Location='LEGACY 101'))
    client = create_app(visit_db_path=str(db)).test_client()
    assert client.get('/api/rooms/all').json[0]['building'] == 'LEGACY'


def test_xlsx_expanded_size_limit_checked_before_openpyxl(tmp_path, monkeypatch):
    path = tmp_path / 'bomb.xlsx'
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('large.xml', b'x' * 2048)
    monkeypatch.setattr(parser, 'MAX_XLSX_EXPANDED_BYTES', 1024)
    with patch('openpyxl.load_workbook') as workbook:
        with pytest.raises(ValueError, match='expanded|large|limit'):
            parser.load_schedule(path)
        workbook.assert_not_called()


def test_xlsx_understated_dimensions_cannot_hide_malformed_rows(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(['Term', 'Days', 'Times', 'Location', 'Max', 'Delivery Mode'])
    ws.append([202690, 'M', '9:00 AM - 10:00 AM', 'TEST 101', 20, 'Face-to-Face'])
    ws.append([202690, 'M', 'bad', 'TEST 202', 20, 'Face-to-Face'])
    original = io.BytesIO()
    wb.save(original)
    path = tmp_path / 'understated.xlsx'
    with zipfile.ZipFile(io.BytesIO(original.getvalue())) as source, zipfile.ZipFile(path, 'w') as output:
        for member in source.infolist():
            contents = source.read(member.filename)
            if member.filename == 'xl/worksheets/sheet1.xml':
                contents = contents.replace(b'ref="A1:F3"', b'ref="A1:F2"')
            output.writestr(member, contents)
    with pytest.raises(ValueError, match='row 3.*Times'):
        parser.load_schedule(path)


def test_wide_data_row_rejected_before_scanning_cells():
    class NoCellAccess:
        def __len__(self):
            return parser.MAX_SCHEDULE_COLUMNS + 1
        def __iter__(self):
            raise AssertionError('oversized row must not be scanned')
        def __getitem__(self, key):
            raise AssertionError('oversized row must not be sliced')
    with pytest.raises(parser.ScheduleValidationError, match='column.*limit'):
        parser._row(['Days', 'Times', 'Location', 'Delivery Mode'], NoCellAccess(), 3, sparse=True)


def test_actual_xlsx_width_limit_survives_understated_dimensions(tmp_path, monkeypatch):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['Term', 'Days', 'Times', 'Location', 'Max', 'Delivery Mode'])
    sheet.append([202690, 'M', '9:00 AM - 10:00 AM', 'TEST 101', 20, 'Face-to-Face'])
    stream = io.BytesIO()
    workbook.save(stream)
    path = tmp_path / 'wide-understated.xlsx'
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source, zipfile.ZipFile(path, 'w') as output:
        for member in source.infolist():
            contents = source.read(member.filename)
            if member.filename == 'xl/worksheets/sheet1.xml':
                extra = b''.join(f'<row r="{number}"><c r="XFD{number}"/></row>'.encode() for number in range(3, 6))
                contents = contents.replace(b'</sheetData>', extra + b'</sheetData>')
            output.writestr(member, contents)
    seen = []
    original_row = parser._row
    def observed_row(headers, values, number, **kwargs):
        seen.append(number)
        return original_row(headers, values, number, **kwargs)
    monkeypatch.setattr(parser, '_row', observed_row)
    with pytest.raises(parser.ScheduleValidationError, match='row 3.*column.*limit'):
        parser.load_schedule(path)
    assert seen == [2, 3]  # stop at the first wide row, never scan subsequent rows


def test_malformed_xlsx_xml_has_bounded_validation_error(tmp_path):
    workbook = Workbook()
    workbook.active.append(['Days', 'Times', 'Location', 'Delivery Mode'])
    stream = io.BytesIO()
    workbook.save(stream)
    path = tmp_path / 'broken.xml.xlsx'
    with zipfile.ZipFile(io.BytesIO(stream.getvalue())) as source, zipfile.ZipFile(path, 'w') as output:
        for member in source.infolist():
            contents = source.read(member.filename)
            if member.filename == 'xl/worksheets/sheet1.xml':
                contents = b'<worksheet><untrusted-text'
            output.writestr(member, contents)
    with pytest.raises(parser.ScheduleValidationError) as error:
        parser.load_schedule(path)
    assert len(str(error.value)) < 2500
    assert 'untrusted-text' not in str(error.value)


def test_parser_row_limit_is_explicit(tmp_path, monkeypatch):
    path = tmp_path / 'many.csv'
    body = csv_bytes()
    path.write_bytes(body + body.split(b'\r\n', 1)[1] * 3)
    monkeypatch.setattr(parser, 'MAX_SCHEDULE_ROWS', 2)
    with pytest.raises(ValueError, match='row.*limit|limit.*row'):
        parser.load_schedule(path)


def test_duplicate_headers_and_extra_csv_cells_rejected(tmp_path):
    for body in (b'Days,Days,Times,Location,Delivery Mode\nM,T,9:00 AM - 10:00 AM,TEST 101,Face-to-Face\n',
                 csv_bytes().replace(b'Face-to-Face\r\n', b'Face-to-Face,unexpected\r\n'),
                 csv_bytes() + b'202690,PARTIAL ROW\r\n'):
        path = tmp_path / 'invalid.csv'
        path.write_bytes(body)
        with pytest.raises(ValueError):
            parser.load_schedule(path)
